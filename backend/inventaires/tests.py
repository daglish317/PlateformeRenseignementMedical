from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from stock.models import Medicament, StockItem, StockMovement
from structures.models import (
    EquipeStructure,
    RoleEquipeStructure,
    StatutEquipeStructure,
    StatutStructure,
    Structure,
    TypeStructure,
)
from utilisateurs.models import RoleUtilisateur, TypeAuthentification, Utilisateur

from .exports import generer_inventaire_excel, generer_inventaire_pdf
from .models import Inventaire, InventaireLigne, StatutInventaire
from .services import InventaireService


class InventaireBaseTestCase(TestCase):

    def setUp(self):
        self.client = APIClient()

        self.structure = Structure.objects.create(
            nom="Pharmacie Test",
            type=TypeStructure.PHARMACIE,
            statut=StatutStructure.ACTIVE,
            adresse="Abidjan",
            telephone="0102030405",
        )

        self.gestionnaire = Utilisateur.objects.create_user(
            email="gestionnaire@test.com",
            password="password123",
            nom="Gest",
            role=RoleUtilisateur.GESTIONNAIRE,
            type_authentification=TypeAuthentification.EMAIL,
            email_verifie=True,
        )
        self.structure.gestionnaire = self.gestionnaire
        self.structure.save(update_fields=["gestionnaire"])
        EquipeStructure.objects.create(
            structure=self.structure,
            utilisateur=self.gestionnaire,
            role=RoleEquipeStructure.GESTIONNAIRE,
            statut=StatutEquipeStructure.ACTIF,
        )

        self.caissier = Utilisateur.objects.create_user(
            email="caissier@test.com",
            password="password123",
            nom="Cais",
            role=RoleUtilisateur.CAISSIER,
            type_authentification=TypeAuthentification.EMAIL,
            email_verifie=True,
        )
        EquipeStructure.objects.create(
            structure=self.structure,
            utilisateur=self.caissier,
            role=RoleEquipeStructure.CAISSIER,
            statut=StatutEquipeStructure.ACTIF,
        )

        self.proprietaire = Utilisateur.objects.create_user(
            email="proprietaire@test.com",
            password="password123",
            nom="Proprio",
            role=RoleUtilisateur.PROPRIETAIRE,
            type_authentification=TypeAuthentification.EMAIL,
            email_verifie=True,
        )
        EquipeStructure.objects.create(
            structure=self.structure,
            utilisateur=self.proprietaire,
            role=RoleEquipeStructure.PROPRIETAIRE,
            statut=StatutEquipeStructure.ACTIF,
        )

        self.gestionnaire_autre = Utilisateur.objects.create_user(
            email="gest-autre@test.com",
            password="password123",
            nom="Gest Autre",
            role=RoleUtilisateur.GESTIONNAIRE,
            type_authentification=TypeAuthentification.EMAIL,
            email_verifie=True,
        )

        self.medicament_disponible = self._medicament(
            "Paracétamol 500mg",
            "COMPRIME",
            quantite=100,
            quantite_reservee=0,
        )
        self.medicament_stock_faible = self._medicament(
            "Amoxicilline 250mg",
            "GELULE",
            quantite=3,
            quantite_reservee=0,
        )
        self.medicament_rupture = self._medicament(
            "Sirop Toux",
            "SIROP",
            quantite=0,
            quantite_reservee=0,
        )
        self.medicament_sans_stock = Medicament.objects.create(
            structure=self.structure,
            nom="Vitamine C",
            forme_pharmaceutique="COMPRIME",
        )
        self.medicament_reserve = self._medicament(
            "Antihistamine",
            "SACHET",
            quantite=10,
            quantite_reservee=3,
        )

        StockItem.objects.create(
            structure=self.structure,
            nom="Tensiomètre",
            type_item=StockItem.TYPE_EQUIPEMENT,
            quantite=5,
        )

    def _medicament(self, nom, forme, quantite, quantite_reservee):
        medicament = Medicament.objects.create(
            structure=self.structure,
            nom=nom,
            forme_pharmaceutique=forme,
        )
        StockItem.objects.create(
            structure=self.structure,
            nom=nom,
            type_item=StockItem.TYPE_MEDICAMENT,
            quantite=quantite,
            quantite_reservee=quantite_reservee,
        )
        return medicament

    def _generer(self, user=None):
        self.client.force_authenticate(user=user or self.gestionnaire)
        return self.client.post(
            "/api/inventaires/generer/",
            {"structure_id": str(self.structure.id)},
        )

    def _inventaire(self):
        return Inventaire.objects.get(
            structure=self.structure,
            cree_par=self.gestionnaire,
        )


class GenerationInventaireTests(InventaireBaseTestCase):

    def test_gestionnaire_genere_un_inventaire(self):
        response = self._generer()
        self.assertEqual(response.status_code, 201)
        self.assertIn("lignes", response.data)
        self.assertTrue(response.data["numero"].startswith("INV"))
        self.assertEqual(response.data["structure_nom"], "Pharmacie Test")
        self.assertEqual(response.data["cree_par_nom"], "Gest")
        self.assertEqual(response.data["role_createur"], "GESTIONNAIRE")
        self.assertEqual(response.data["nombre_total_produits"], 5)
        self.assertEqual(len(response.data["lignes"]), 5)

    def test_numero_sequentiel_par_structure(self):
        self._generer()
        self._generer()
        numeros = list(
            Inventaire.objects.filter(
                structure=self.structure,
            ).values_list("numero", flat=True)
        )
        self.assertEqual(len(numeros), 2)
        self.assertNotEqual(numeros[0], numeros[1])

    def test_lignes_captent_les_donnees_du_stock(self):
        self._generer()
        ligne = InventaireLigne.objects.get(
            inventaire=self._inventaire(),
            nom="Paracétamol 500mg",
        )
        self.assertEqual(ligne.quantite_physique, 100)
        self.assertEqual(ligne.quantite_reservee, 0)
        self.assertEqual(ligne.quantite_disponible, 100)
        self.assertEqual(ligne.seuil_alerte, 5)
        self.assertEqual(ligne.statut, StatutInventaire.DISPONIBLE)
        self.assertEqual(ligne.forme_pharmaceutique, "Comprimé")

    def test_inventaire_ne_modifie_pas_le_stock(self):
        avant = {
            item.id: (item.quantite, item.quantite_reservee)
            for item in StockItem.objects.filter(
                structure=self.structure,
                type_item=StockItem.TYPE_MEDICAMENT,
            )
        }
        self._generer()
        apres = {
            item.id: (item.quantite, item.quantite_reservee)
            for item in StockItem.objects.filter(
                structure=self.structure,
                type_item=StockItem.TYPE_MEDICAMENT,
            )
        }
        self.assertEqual(avant, apres)
        self.assertEqual(StockMovement.objects.count(), 0)

    def test_ancien_inventaire_immuable(self):
        self._generer()
        ligne = InventaireLigne.objects.get(
            inventaire=self._inventaire(),
            nom="Paracétamol 500mg",
        )
        item = StockItem.objects.get(
            structure=self.structure,
            nom="Paracétamol 500mg",
            type_item=StockItem.TYPE_MEDICAMENT,
        )
        item.quantite = 10
        item.save()

        ligne.refresh_from_db()
        self.assertEqual(ligne.quantite_physique, 100)
        self.assertEqual(ligne.quantite_disponible, 100)

    def test_statuts_calcules(self):
        self._generer()
        lignes = {
            ligne.nom: ligne.statut
            for ligne in self._inventaire().lignes.all()
        }
        self.assertEqual(
            lignes["Paracétamol 500mg"],
            StatutInventaire.DISPONIBLE,
        )
        self.assertEqual(
            lignes["Antihistamine"],
            StatutInventaire.DISPONIBLE,
        )
        self.assertEqual(
            lignes["Amoxicilline 250mg"],
            StatutInventaire.STOCK_FAIBLE,
        )
        self.assertEqual(
            lignes["Sirop Toux"],
            StatutInventaire.RUPTURE,
        )
        self.assertEqual(
            lignes["Vitamine C"],
            StatutInventaire.RUPTURE,
        )

    def test_resume_general(self):
        response = self._generer()
        self.assertEqual(response.data["nombre_disponibles"], 2)
        self.assertEqual(response.data["nombre_stock_faible"], 1)
        self.assertEqual(response.data["nombre_ruptures"], 2)

    def test_quantite_disponible_soustrait_la_reservee(self):
        self._generer()
        ligne = InventaireLigne.objects.get(
            inventaire=self._inventaire(),
            nom="Antihistamine",
        )
        self.assertEqual(ligne.quantite_physique, 10)
        self.assertEqual(ligne.quantite_reservee, 3)
        self.assertEqual(ligne.quantite_disponible, 7)

    def test_equipement_exclu_de_l_inventaire(self):
        self._generer()
        noms = list(self._inventaire().lignes.values_list("nom", flat=True))
        self.assertNotIn("Tensiomètre", noms)

    def test_structure_id_requis(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.post("/api/inventaires/generer/", {})
        self.assertEqual(response.status_code, 400)

    def test_medicament_sans_stock_est_une_rupture(self):
        self._generer()
        ligne = InventaireLigne.objects.get(
            inventaire=self._inventaire(),
            nom="Vitamine C",
        )
        self.assertEqual(ligne.quantite_physique, 0)
        self.assertEqual(ligne.quantite_disponible, 0)
        self.assertEqual(ligne.statut, StatutInventaire.RUPTURE)


class ConsultationInventaireTests(InventaireBaseTestCase):

    def setUp(self):
        super().setUp()
        self._generer()
        self.inventaire = self._inventaire()

    def test_liste_inventaires(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            "/api/inventaires/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]["id"], str(self.inventaire.id))
        self.assertEqual(response.data[0]["nombre_total_produits"], 5)

    def test_detail_avec_lignes(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(f"/api/inventaires/{self.inventaire.id}/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data["lignes"]), 5)
        premiere = response.data["lignes"][0]
        self.assertIn("statut_label", premiere)

    def test_recherche_par_nom(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            f"/api/inventaires/{self.inventaire.id}/",
            {"recherche": "paracétamol"},
        )
        lignes = response.data["lignes"]
        self.assertEqual(len(lignes), 1)
        self.assertEqual(lignes[0]["nom"], "Paracétamol 500mg")

    def test_recherche_par_forme(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            f"/api/inventaires/{self.inventaire.id}/",
            {"recherche": "gélule"},
        )
        lignes = response.data["lignes"]
        self.assertEqual(len(lignes), 1)
        self.assertEqual(lignes[0]["nom"], "Amoxicilline 250mg")

    def test_recherche_par_statut_libelle(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            f"/api/inventaires/{self.inventaire.id}/",
            {"recherche": "rupture"},
        )
        lignes = response.data["lignes"]
        self.assertEqual(len(lignes), 2)

    def test_filtre_par_statut(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            f"/api/inventaires/{self.inventaire.id}/",
            {"statut": "STOCK_FAIBLE"},
        )
        lignes = response.data["lignes"]
        self.assertEqual(len(lignes), 1)
        self.assertEqual(lignes[0]["statut"], "STOCK_FAIBLE")


class PermissionsInventaireTests(InventaireBaseTestCase):

    def setUp(self):
        super().setUp()
        self._generer()
        self.inventaire = self._inventaire()

    def test_proprietaire_consulte_la_liste(self):
        self.client.force_authenticate(user=self.proprietaire)
        response = self.client.get(
            "/api/inventaires/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.data), 1)

    def test_proprietaire_consulte_le_detail(self):
        self.client.force_authenticate(user=self.proprietaire)
        response = self.client.get(f"/api/inventaires/{self.inventaire.id}/")
        self.assertEqual(response.status_code, 200)

    def test_proprietaire_ne_peut_pas_generer(self):
        self.client.force_authenticate(user=self.proprietaire)
        response = self.client.post(
            "/api/inventaires/generer/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 403)

    def test_caissier_n_a_aucun_acces(self):
        self.client.force_authenticate(user=self.caissier)
        response = self.client.get(
            "/api/inventaires/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 403)

        response = self.client.post(
            "/api/inventaires/generer/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 403)

        response = self.client.get(f"/api/inventaires/{self.inventaire.id}/")
        self.assertEqual(response.status_code, 403)

    def test_gestionnaire_autre_structure_refuse(self):
        self.client.force_authenticate(user=self.gestionnaire_autre)
        response = self.client.get(
            "/api/inventaires/",
            {"structure_id": str(self.structure.id)},
        )
        self.assertEqual(response.status_code, 403)

        response = self.client.get(f"/api/inventaires/{self.inventaire.id}/")
        self.assertEqual(response.status_code, 403)

    def test_non_authentifie_refuse(self):
        self.client.force_authenticate(user=None)
        response = self.client.get("/api/inventaires/")
        self.assertEqual(response.status_code, 401)

    def test_inventaire_introuvable(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            "/api/inventaires/00000000-0000-0000-0000-000000000000/"
        )
        self.assertEqual(response.status_code, 404)


class ExportInventaireTests(InventaireBaseTestCase):

    def setUp(self):
        super().setUp()
        self._generer()
        self.inventaire = self._inventaire()

    def test_pdf_telechargeable(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(f"/api/inventaires/{self.inventaire.id}/pdf/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

        contenu = generer_inventaire_pdf(self.inventaire)
        self.assertTrue(contenu.startswith(b"%PDF"))

    def test_excel_telechargeable(self):
        self.client.force_authenticate(user=self.gestionnaire)
        response = self.client.get(
            f"/api/inventaires/{self.inventaire.id}/excel/"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            response["Content-Type"],
        )

        wb = load_workbook(BytesIO(generer_inventaire_excel(self.inventaire)))
        ws = wb.active
        valeurs = [
            str(cell.value)
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        self.assertIn("Inventaire n°", " ".join(valeurs))
        self.assertIn("Paracétamol 500mg", " ".join(valeurs))
        self.assertIn("Amoxicilline 250mg", " ".join(valeurs))
        self.assertIn("Disponible", " ".join(valeurs))

    def test_proprietaire_peut_exporter(self):
        self.client.force_authenticate(user=self.proprietaire)
        pdf = self.client.get(f"/api/inventaires/{self.inventaire.id}/pdf/")
        self.assertEqual(pdf.status_code, 200)
        excel = self.client.get(f"/api/inventaires/{self.inventaire.id}/excel/")
        self.assertEqual(excel.status_code, 200)

    def test_caissier_ne_peut_pas_exporter(self):
        self.client.force_authenticate(user=self.caissier)
        pdf = self.client.get(f"/api/inventaires/{self.inventaire.id}/pdf/")
        self.assertEqual(pdf.status_code, 403)
        excel = self.client.get(f"/api/inventaires/{self.inventaire.id}/excel/")
        self.assertEqual(excel.status_code, 403)


class ServiceInventaireTests(InventaireBaseTestCase):

    def test_statut_medicament(self):
        self.assertEqual(
            InventaireService.statut_medicament(0, 5),
            StatutInventaire.RUPTURE,
        )
        self.assertEqual(
            InventaireService.statut_medicament(5, 5),
            StatutInventaire.STOCK_FAIBLE,
        )
        self.assertEqual(
            InventaireService.statut_medicament(3, 5),
            StatutInventaire.STOCK_FAIBLE,
        )
        self.assertEqual(
            InventaireService.statut_medicament(6, 5),
            StatutInventaire.DISPONIBLE,
        )

    def test_generation_directe_via_service(self):
        inventaire = InventaireService.generer(
            structure=self.structure,
            utilisateur=self.gestionnaire,
        )
        self.assertEqual(inventaire.cree_par, self.gestionnaire)
        self.assertEqual(inventaire.lignes.count(), 5)
