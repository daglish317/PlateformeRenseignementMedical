"""Exports PDF et Excel d'un inventaire (module Inventaire)."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from django.db.models import Count, Q
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .models import StatutInventaire

_PDF_STYLE_FOND = colors.HexColor("#e2e8f0")
_PDF_STYLE_BORDURE = colors.HexColor("#cbd5e1")
_PDF_STYLE_LIGNE = colors.HexColor("#334155")

_XLSX_STYLE_FOND = PatternFill(
    start_color="E2E8F0",
    end_color="E2E8F0",
    fill_type="solid",
)
_XLSX_STYLE_BORDURE = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _resume(inventaire):
    """Résumé général conforme à la spécification (niveau 1)."""
    compteurs = inventaire.lignes.aggregate(
        total=Count("id"),
        disponibles=Count("id", filter=Q(statut=StatutInventaire.DISPONIBLE)),
        stock_faible=Count("id", filter=Q(statut=StatutInventaire.STOCK_FAIBLE)),
        ruptures=Count("id", filter=Q(statut=StatutInventaire.RUPTURE)),
    )
    return compteurs


def _lignes_document(inventaire, lignes=None):
    if lignes is None:
        lignes = inventaire.lignes.all()
    return lignes


def _titre_style(nom, taille=16, gras=True, align=1):
    return ParagraphStyle(
        name=nom,
        fontName="Helvetica-Bold" if gras else "Helvetica",
        fontSize=taille,
        leading=taille + 4,
        spaceAfter=2,
        alignment=align,
    )


def generer_inventaire_pdf(inventaire, lignes=None):
    """Retourne les octets du PDF d'état du stock (document historique)."""
    lignes = _lignes_document(inventaire, lignes)
    resume = _resume(inventaire)
    structure = inventaire.structure

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"Inventaire {inventaire.numero}",
    )

    elements = []
    elements.append(Paragraph(structure.nom or "Pharmacie", _titre_style("Titre")))
    elements.append(Paragraph(
        f"Inventaire n° {inventaire.numero}",
        _titre_style("SousTitre", taille=12, gras=False),
    ))
    elements.append(Spacer(1, 6 * mm))

    infos = Table(
        [
            ["Date", f"{inventaire.cree_le:%d/%m/%Y}"],
            ["Heure", f"{inventaire.cree_le:%H:%M}"],
            ["Généré par", inventaire.cree_par.nom],
            ["Rôle", inventaire.role_createur or "—"],
        ],
        colWidths=[45 * mm, None],
    )
    infos.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(infos)
    elements.append(Spacer(1, 6 * mm))

    resume_data = [
        [
            "Total produits",
            "Disponibles",
            "Stock faible",
            "Ruptures",
        ],
        [
            str(resume["total"]),
            str(resume["disponibles"]),
            str(resume["stock_faible"]),
            str(resume["ruptures"]),
        ],
    ]
    resume_table = Table(
        resume_data,
        colWidths=[45 * mm, 45 * mm, 45 * mm, 45 * mm],
    )
    resume_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_STYLE_FOND),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_STYLE_BORDURE),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(resume_table)
    elements.append(Spacer(1, 6 * mm))

    data = [[
        "Médicament",
        "Forme",
        "Qté physique",
        "Qté réservée",
        "Qté disponible",
        "Seuil",
        "Statut",
    ]]
    for ligne in lignes:
        data.append([
            ligne.nom,
            ligne.forme_pharmaceutique or "—",
            str(ligne.quantite_physique),
            str(ligne.quantite_reservee),
            str(ligne.quantite_disponible),
            str(ligne.seuil_alerte),
            ligne.get_statut_display(),
        ])

    tableau = Table(
        data,
        colWidths=[None, 45 * mm, 32 * mm, 32 * mm, 32 * mm, 25 * mm, 35 * mm],
        repeatRows=1,
    )
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_STYLE_FOND),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (2, 0), (5, -1), "RIGHT"),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_STYLE_BORDURE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(tableau)
    elements.append(Spacer(1, 8 * mm))

    elements.append(Paragraph(
        "Document historique : état du stock au moment de la génération. "
        "L'inventaire n'effectue aucune opération de gestion du stock.",
        _titre_style("Note", taille=8, gras=False, align=0),
    ))

    doc.build(elements)
    return buffer.getvalue()


def generer_inventaire_excel(inventaire, lignes=None):
    """Retourne les octets du classeur Excel d'un inventaire."""
    lignes = _lignes_document(inventaire, lignes)
    resume = _resume(inventaire)
    structure = inventaire.structure

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"

    ws["A1"] = structure.nom or "Pharmacie"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Inventaire n° {inventaire.numero}"

    infos = [
        ("Date", f"{inventaire.cree_le:%d/%m/%Y}"),
        ("Heure", f"{inventaire.cree_le:%H:%M}"),
        ("Généré par", inventaire.cree_par.nom),
        ("Rôle", inventaire.role_createur or "—"),
        (
            "Résumé",
            (
                f"Total : {resume['total']} | Disponibles : "
                f"{resume['disponibles']} | Stock faible : "
                f"{resume['stock_faible']} | Ruptures : "
                f"{resume['ruptures']}"
            ),
        ),
    ]
    for i, (cle, valeur) in enumerate(infos, start=4):
        ws.cell(row=i, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valeur)

    entete_row = 10
    entetes = [
        "Médicament",
        "Forme",
        "Qté physique",
        "Qté réservée",
        "Qté disponible",
        "Seuil",
        "Statut",
    ]
    for col, entete in enumerate(entetes, start=1):
        cellule = ws.cell(row=entete_row, column=col, value=entete)
        cellule.font = Font(bold=True)
        cellule.fill = _XLSX_STYLE_FOND
        cellule.border = _XLSX_STYLE_BORDURE
        cellule.alignment = Alignment(horizontal="center")

    for i, ligne in enumerate(lignes, start=entete_row + 1):
        valeurs = [
            ligne.nom,
            ligne.forme_pharmaceutique or "—",
            ligne.quantite_physique,
            ligne.quantite_reservee,
            ligne.quantite_disponible,
            ligne.seuil_alerte,
            ligne.get_statut_display(),
        ]
        for col, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=i, column=col, value=valeur)
            cellule.border = _XLSX_STYLE_BORDURE

    largeurs = {
        "A": 45,
        "B": 25,
        "C": 14,
        "D": 14,
        "E": 14,
        "F": 10,
        "G": 18,
    }
    for colonne, largeur in largeurs.items():
        ws.column_dimensions[colonne].width = largeur

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
