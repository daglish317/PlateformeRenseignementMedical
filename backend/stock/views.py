import csv
import io
from io import BytesIO

from django.db.models import Exists, F, OuterRef, Subquery
from django.http import FileResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from .models import Approvisionnement, Medicament, StockItem, StockMovement
from .serializers import (
    ApprovisionnementSerializer,
    ApprovisionnementCreateSerializer,
    MedicamentSerializer,
    ProduitPeremptionSerializer,
    StockSerializer,
    StockMovementSerializer,
)
from .exports import generer_approvisionnement_excel, generer_approvisionnement_pdf
from .services import StockService
from historique.models import TypeEvenementHistorique
from historique.services import HistoriqueService

from structures.models import Structure
from structures.permissions import (
    assert_operational_access,
    assert_structure_autorise_stock_direct,
)
from structures.permission_registry import ActionPermission, ModuleOperationnel
from utilisateurs.decorators import operational_member_required


def _parse_csv(file):
    decoded = file.read().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(decoded))
    rows = []
    for row in reader:
        rows.append(dict(row))
    return rows


def _parse_excel(file):
    try:
        import openpyxl
    except ImportError:
        return None

    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    headers = [
        str(cell.value).strip().lower() if cell.value else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]

    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for header, value in zip(headers, row):
            if header:
                row_dict[header] = str(value) if value is not None else ""
        rows.append(row_dict)

    wb.close()
    return rows


class CreateStockView(APIView):
    """Creation directe d'un article de stock (interdite pour les pharmacies)."""

    @operational_member_required
    def post(self, request):

        structure_id = request.data.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user,
            structure_id,
            ModuleOperationnel.STOCK,
            ActionPermission.MODIFIER,
        )
        assert_structure_autorise_stock_direct(request.user, structure_id)
        try:
            structure = Structure.objects.get(id=structure_id)
        except (Structure.DoesNotExist, ValueError):
            return Response(
                {"detail": "Structure introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        nom = request.data.get("nom")
        if not nom:
            return Response({"detail": "Nom manquant"}, status=status.HTTP_400_BAD_REQUEST)

        quantite_raw = request.data.get("quantite", 0)
        try:
            quantite = int(float(quantite_raw))
        except (ValueError, TypeError):
            quantite = 0

        item = StockService.ajouter_ou_mettre_a_jour(
            structure=structure,
            nom=nom,
            type_item=request.data.get("type_item", StockItem.TYPE_MEDICAMENT),
            quantite=quantite,
            disponible=request.data.get("disponible", True),
            seuil_alerte=int(request.data.get("seuil_alerte", 5)),
        )

        return Response(StockSerializer(item).data, status=status.HTTP_201_CREATED)


class ListStockStructureView(APIView):

    @operational_member_required
    def get(self, request, structure_id):

        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.STOCK, ActionPermission.CONSULTER
        )
        items = StockItem.objects.filter(structure_id=structure_id).order_by("nom")
        return Response(StockSerializer(items, many=True).data)


class DeleteStockView(APIView):

    @operational_member_required
    def delete(self, request, pk):

        try:
            item = StockItem.objects.select_related("structure").get(id=pk)
        except StockItem.DoesNotExist:
            return Response(
                {"detail": "Article introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_operational_access(
            request.user, item.structure_id, ModuleOperationnel.STOCK, ActionPermission.SUPPRIMER
        )

        HistoriqueService.enregistrer(
            structure=item.structure,
            type_evenement=TypeEvenementHistorique.STOCK_SUPPRIME,
            utilisateur=request.user,
            donnees={
                "nom": item.nom,
                "quantite": item.quantite,
                "type_item": item.type_item,
            },
        )
        StockService.supprimer_item(item=item)

        return Response(
            {"message": "Article supprimé"},
            status=status.HTTP_200_OK,
        )


class RetirerStockView(APIView):

    @operational_member_required
    def post(self, request, pk):

        try:
            item = StockItem.objects.select_related("structure").get(id=pk)
        except StockItem.DoesNotExist:
            return Response(
                {"detail": "Article introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )
        assert_operational_access(
            request.user, item.structure_id, ModuleOperationnel.STOCK, ActionPermission.MODIFIER
        )
        try:
            item = StockService.retirer_stock(
                item=item,
                quantite=int(request.data.get("quantite", 0)),
                motif=request.data.get("motif", ""),
            )
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(StockSerializer(item).data)


class StockMovementsView(APIView):

    @operational_member_required
    def get(self, request, pk):

        try:
            item = StockItem.objects.select_related("structure").get(id=pk)
        except StockItem.DoesNotExist:
            return Response(
                {"detail": "Article introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )
        assert_operational_access(
            request.user, item.structure_id, ModuleOperationnel.STOCK, ActionPermission.CONSULTER
        )
        mouvements = StockMovement.objects.filter(item=item)
        return Response(StockMovementSerializer(mouvements, many=True).data)


class StockAlertesView(APIView):

    @operational_member_required
    def get(self, request, structure_id):

        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.STOCK, ActionPermission.CONSULTER
        )
        structure = Structure.objects.get(id=structure_id)
        items = StockService.items_en_alerte(structure).order_by("nom")
        return Response(StockSerializer(items, many=True).data)


class ProduitsPeremptionView(APIView):

    @operational_member_required
    def get(self, request, structure_id):

        from datetime import timedelta
        from .models import LigneApprovisionnement

        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.PEREMPTION, ActionPermission.CONSULTER
        )
        limite = timezone.localdate() + timedelta(days=92)
        lignes = (
            LigneApprovisionnement.objects.filter(
                approvisionnement__structure_id=structure_id,
                date_peremption__lte=limite,
            )
            .select_related("approvisionnement", "medicament")
            .order_by("date_peremption", "medicament__nom")
        )
        return Response(ProduitPeremptionSerializer(lignes, many=True).data)


class MedicamentListView(APIView):
    """Recherche d'autocomplétion pour le formulaire d'approvisionnement."""

    @operational_member_required
    def get(self, request):

        structure_id = request.query_params.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.RECHERCHER
        )

        queryset = Medicament.objects.filter(structure_id=structure_id)
        vendable = (
            (request.query_params.get("vendable") or "").strip().lower()
            in {"1", "true", "oui", "yes"}
        )
        if vendable:
            stock_disponible = StockItem.objects.filter(
                structure_id=structure_id,
                nom=OuterRef("nom"),
                quantite__gt=F("quantite_reservee"),
            )
            queryset = (
                queryset.filter(
                    prix_vente__gt=0,
                    en_reserve=False,
                )
                .annotate(a_stock_disponible=Exists(stock_disponible))
                .filter(a_stock_disponible=True)
            )
        search = (request.query_params.get("search") or "").strip()
        if search:
            queryset = queryset.filter(nom__icontains=search)

        stock_item = StockItem.objects.filter(
            structure_id=OuterRef("structure_id"),
            nom=OuterRef("nom"),
        )
        queryset = queryset.annotate(
            _stock_avant=Subquery(stock_item.values("quantite")[:1]),
            _stock_physique=Subquery(stock_item.values("quantite")[:1]),
            _stock_reservee=Subquery(
                stock_item.values("quantite_reservee")[:1]
            ),
        )

        queryset = queryset.order_by("nom")[:20]
        return Response(MedicamentSerializer(queryset, many=True).data)


class ApprovisionnementListCreateView(APIView):

    @operational_member_required
    def get(self, request):

        structure_id = request.query_params.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.CONSULTER
        )

        queryset = (
            Approvisionnement.objects.filter(structure_id=structure_id)
            .select_related("structure", "cree_par")
            .prefetch_related("lignes__medicament")
        )
        return Response(ApprovisionnementSerializer(queryset, many=True).data)

    @operational_member_required
    def post(self, request):

        structure_id = request.data.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.CREER
        )
        structure = Structure.objects.get(id=structure_id)

        serializer = ApprovisionnementCreateSerializer(
            data=request.data,
            context={"request": request, "structure": structure},
        )
        serializer.is_valid(raise_exception=True)
        try:
            appro = serializer.save()
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(
            ApprovisionnementSerializer(appro).data,
            status=status.HTTP_201_CREATED,
        )


class ApprovisionnementDetailView(APIView):

    @operational_member_required
    def get(self, request, pk):

        try:
            appro = (
                Approvisionnement.objects.select_related("structure", "cree_par")
                .prefetch_related("lignes__medicament")
                .get(id=pk)
            )
        except Approvisionnement.DoesNotExist:
            return Response(
                {"detail": "Approvisionnement introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_operational_access(
            request.user, appro.structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.CONSULTER
        )
        return Response(ApprovisionnementSerializer(appro).data)


class ApprovisionnementPDFView(APIView):

    @operational_member_required
    def get(self, request, pk):
        try:
            appro = (
                Approvisionnement.objects.select_related("structure", "cree_par")
                .prefetch_related("lignes__medicament")
                .get(id=pk)
            )
        except Approvisionnement.DoesNotExist:
            return Response(
                {"detail": "Approvisionnement introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_operational_access(
            request.user, appro.structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.IMPRIMER
        )
        pdf = BytesIO(generer_approvisionnement_pdf(appro))
        nom = f"approvisionnement_{appro.numero}.pdf".replace(" ", "_")
        return FileResponse(
            pdf,
            as_attachment=True,
            filename=nom,
            content_type="application/pdf",
        )


class ApprovisionnementExcelView(APIView):

    @operational_member_required
    def get(self, request, pk):
        try:
            appro = (
                Approvisionnement.objects.select_related("structure", "cree_par")
                .prefetch_related("lignes__medicament")
                .get(id=pk)
            )
        except Approvisionnement.DoesNotExist:
            return Response(
                {"detail": "Approvisionnement introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_operational_access(
            request.user, appro.structure_id, ModuleOperationnel.APPROVISIONNEMENT, ActionPermission.EXPORTER
        )
        excel = BytesIO(generer_approvisionnement_excel(appro))
        nom = f"approvisionnement_{appro.numero}.xlsx".replace(" ", "_")
        return FileResponse(
            excel,
            as_attachment=True,
            filename=nom,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )


class EntreeStockView(APIView):
    """Entree de stock directe (interdite pour les pharmacies)."""

    @operational_member_required
    def post(self, request, pk):

        try:
            item = StockItem.objects.select_related("structure").get(id=pk)
        except StockItem.DoesNotExist:
            return Response(
                {"detail": "Article introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )
        assert_operational_access(
            request.user, item.structure_id, ModuleOperationnel.STOCK, ActionPermission.MODIFIER
        )
        assert_structure_autorise_stock_direct(request.user, item.structure_id)
        try:
            item = StockService.entree_stock(
                item=item,
                quantite=int(request.data.get("quantite", 0)),
                motif=request.data.get("motif", ""),
            )
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(StockSerializer(item).data)


class ImportMedicamentView(APIView):
    """Import de medicaments via fichier (interdit pour les pharmacies)."""

    @operational_member_required
    def post(self, request):
        structure_id = request.data.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.STOCK, ActionPermission.MODIFIER
        )
        assert_structure_autorise_stock_direct(request.user, structure_id)
        try:
            structure = Structure.objects.get(id=structure_id)
        except (Structure.DoesNotExist, ValueError):
            return Response(
                {"detail": "Structure introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        file = request.FILES.get("file")
        if not file:
            return Response({"detail": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)

        filename = file.name.lower()
        if filename.endswith(".csv"):
            rows = _parse_csv(file)
        elif filename.endswith((".xlsx", ".xls")):
            rows = _parse_excel(file)
        else:
            return Response({"detail": "Format non supporté. Utilisez CSV ou Excel (.xlsx)"}, status=status.HTTP_400_BAD_REQUEST)

        if not rows:
            return Response({"detail": "Le fichier est vide ou mal formaté"}, status=status.HTTP_400_BAD_REQUEST)

        imported = []
        errors = []

        for i, row in enumerate(rows, start=2):
            try:
                nom = (row.get("nom") or "").strip()
                if not nom:
                    errors.append({"ligne": i, "erreur": "Nom manquant"})
                    continue

                quantite_raw = row.get("quantite") or "0"
                try:
                    quantite = int(float(quantite_raw))
                except (ValueError, TypeError):
                    quantite = 0

                if quantite < 0:
                    errors.append({"ligne": i, "erreur": "Quantité négative"})
                    continue

                item = StockService.ajouter_ou_mettre_a_jour(
                    structure=structure,
                    nom=nom,
                    type_item=StockItem.TYPE_MEDICAMENT,
                    quantite=quantite,
                    disponible=True,
                    seuil_alerte=5,
                )
                imported.append({"id": str(item.id), "nom": item.nom, "type_item": item.type_item})

            except Exception as e:
                errors.append({"ligne": i, "erreur": str(e)})

        return Response({
            "message": f"{len(imported)} médicament(s) importé(s)",
            "imported": len(imported),
            "errors_count": len(errors),
            "errors": errors[:20],
        }, status=status.HTTP_200_OK)


class ImportStockView(APIView):
    """Import generique de stock via fichier (interdit pour les pharmacies)."""

    @operational_member_required
    def post(self, request):
        structure_id = request.data.get("structure_id")
        if not structure_id:
            return Response(
                {"detail": "structure_id requis"},
                status=status.HTTP_400_BAD_REQUEST,
            )
        assert_operational_access(
            request.user, structure_id, ModuleOperationnel.STOCK, ActionPermission.MODIFIER
        )
        assert_structure_autorise_stock_direct(request.user, structure_id)
        try:
            structure = Structure.objects.get(id=structure_id)
        except (Structure.DoesNotExist, ValueError):
            return Response(
                {"detail": "Structure introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        file = request.FILES.get("file")
        if not file:
            return Response(
                {"detail": "Aucun fichier fourni"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        filename = file.name.lower()
        if filename.endswith(".csv"):
            rows = _parse_csv(file)
        elif filename.endswith((".xlsx", ".xls")):
            rows = _parse_excel(file)
        else:
            return Response(
                {"detail": "Format non supporté. Utilisez CSV ou Excel (.xlsx)"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not rows:
            return Response(
                {"detail": "Le fichier est vide ou mal formaté"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        imported = []
        errors = []

        for i, row in enumerate(rows, start=2):
            try:
                nom = row.get("nom", "").strip()
                type_item = row.get("type_item", "").strip().upper()
                quantite = int(row.get("quantite", 0))
                seuil_alerte = int(row.get("seuil_alerte", 5))
                disponible_raw = row.get("disponible", "true").strip().lower()
                disponible = disponible_raw in ("true", "1", "oui", "yes", "vrai")

                if not nom:
                    errors.append({"ligne": i, "erreur": "Nom manquant"})
                    continue

                if type_item not in (
                    StockItem.TYPE_MEDICAMENT,
                    StockItem.TYPE_EQUIPEMENT,
                    StockItem.TYPE_CONSOMMABLE,
                ):
                    errors.append({"ligne": i, "erreur": f"Type invalide: {type_item}"})
                    continue

                if quantite < 0:
                    errors.append({"ligne": i, "erreur": "Quantité négative"})
                    continue

                item = StockService.ajouter_ou_mettre_a_jour(
                    structure=structure,
                    nom=nom,
                    type_item=type_item,
                    quantite=quantite,
                    disponible=disponible,
                    seuil_alerte=seuil_alerte,
                )
                imported.append({"id": str(item.id), "nom": item.nom, "type_item": item.type_item})

            except Exception as e:
                errors.append({"ligne": i, "erreur": str(e)})

        return Response(
            {
                "message": f"{len(imported)} article(s) importé(s)",
                "imported": len(imported),
                "errors_count": len(errors),
                "errors": errors[:20],
            },
            status=status.HTTP_200_OK,
        )
