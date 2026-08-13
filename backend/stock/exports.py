"""Exports PDF et Excel des approvisionnements."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


_HEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _montant_ligne(ligne):
    return ligne.prix_achat * ligne.quantite


def _titre_style(name, size=16, bold=True, align=1):
    return ParagraphStyle(
        name=name,
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=size,
        leading=size + 4,
        alignment=align,
    )


def generer_approvisionnement_excel(approvisionnement):
    wb = Workbook()
    ws = wb.active
    ws.title = "Approvisionnement"

    ws["A1"] = approvisionnement.structure.nom
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Bon d'approvisionnement"

    infos = [
        ("Numero approvisionnement", approvisionnement.numero or ""),
        ("Numero bon de livraison", approvisionnement.reference_bon or ""),
        ("Date approvisionnement", f"{approvisionnement.date_reception:%d/%m/%Y}"),
        ("Fournisseur", approvisionnement.fournisseur or ""),
        ("Montant declare", float(approvisionnement.montant_total_declare)),
        ("Cree par", approvisionnement.cree_par.nom),
    ]
    for row, (label, value) in enumerate(infos, start=4):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)

    header_row = 12
    headers = [
        "Produit",
        "Type",
        "Stock avant",
        "Quantite ajoutee",
        "Stock final attendu",
        "Prix achat",
        "Prix vente",
        "TVA",
        "Reserve",
        "Date peremption",
        "Montant ligne",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center")

    for row, ligne in enumerate(approvisionnement.lignes.all(), start=header_row + 1):
        values = [
            ligne.medicament.nom,
            ligne.get_forme_pharmaceutique_display(),
            ligne.stock_avant,
            ligne.quantite,
            ligne.stock_avant + ligne.quantite,
            float(ligne.prix_achat),
            "" if ligne.prix_vente is None else float(ligne.prix_vente),
            "Oui" if ligne.tva else "Non",
            "Oui" if ligne.en_reserve else "Non",
            f"{ligne.date_peremption:%d/%m/%Y}",
            float(_montant_ligne(ligne)),
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = _BORDER

    for column, width in {
        "A": 28,
        "B": 18,
        "C": 14,
        "D": 18,
        "E": 20,
        "F": 14,
        "G": 14,
        "H": 10,
        "I": 10,
        "J": 16,
        "K": 16,
    }.items():
        ws.column_dimensions[column].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generer_approvisionnement_pdf(approvisionnement):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Approvisionnement {approvisionnement.numero}",
    )

    elements = [
        Paragraph(approvisionnement.structure.nom or "Structure", _titre_style("Title")),
        Paragraph("Bon d'approvisionnement", _titre_style("Subtitle", size=11, bold=False)),
        Spacer(1, 5 * mm),
    ]

    infos = Table(
        [
            ["Numero approvisionnement", approvisionnement.numero or ""],
            ["Numero bon de livraison", approvisionnement.reference_bon or ""],
            ["Date", f"{approvisionnement.date_reception:%d/%m/%Y}"],
            ["Fournisseur", approvisionnement.fournisseur or ""],
            ["Montant declare", f"{approvisionnement.montant_total_declare:,.2f}"],
            ["Cree par", approvisionnement.cree_par.nom],
        ],
        colWidths=[45 * mm, None],
    )
    infos.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(infos)
    elements.append(Spacer(1, 5 * mm))

    data = [[
        "Produit",
        "Type",
        "Stock avant",
        "Ajoutee",
        "Stock final",
        "Prix achat",
        "Prix vente",
        "TVA",
        "Reserve",
        "Peremption",
        "Montant",
    ]]
    for ligne in approvisionnement.lignes.all():
        data.append([
            ligne.medicament.nom,
            ligne.get_forme_pharmaceutique_display(),
            str(ligne.stock_avant),
            str(ligne.quantite),
            str(ligne.stock_avant + ligne.quantite),
            f"{ligne.prix_achat:,.2f}",
            "" if ligne.prix_vente is None else f"{ligne.prix_vente:,.2f}",
            "Oui" if ligne.tva else "Non",
            "Oui" if ligne.en_reserve else "Non",
            f"{ligne.date_peremption:%d/%m/%Y}",
            f"{_montant_ligne(ligne):,.2f}",
        ])

    table = Table(
        data,
        colWidths=[40 * mm, 25 * mm, 22 * mm, 20 * mm, 22 * mm, 24 * mm, 24 * mm, 14 * mm, 16 * mm, 24 * mm, 25 * mm],
        repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)

    doc.build(elements)
    return buffer.getvalue()
