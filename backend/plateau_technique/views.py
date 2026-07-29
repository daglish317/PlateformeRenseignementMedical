import csv
import io

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from .models import PlateauTechnique
from .serializers import PlateauTechniqueSerializer, PlateauTechniqueCreateSerializer
from .services import PlateauTechniqueService

from structures.models import Structure, StructureService
from utilisateurs.decorators import gestionnaire_required, admin_required
from structures.permissions import assert_gestionnaire_owns_structure


class CreatePlateauTechniqueView(APIView):

    @gestionnaire_required
    def post(self, request):

        serializer = PlateauTechniqueCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        assert_gestionnaire_owns_structure(request.user, data["structure_id"])

        if not data.get("service_id") and data.get("nom"):
            service_obj, _ = StructureService.objects.get_or_create(
                nom=data["nom"],
                defaults={"type": "EXAMEN", "description": ""},
            )
            data["service_id"] = service_obj.id

        plateau = PlateauTechniqueService.ajouter_ou_mettre_a_jour(
            structure_id=data["structure_id"],
            service_id=data["service_id"],
            disponible=data["disponible"],
        )

        return Response(
            {"message": "Équipement enregistré", "data": PlateauTechniqueSerializer(plateau).data},
            status=status.HTTP_201_CREATED,
        )


class ListPlateauStructureView(APIView):

    def get(self, request, structure_id):

        data = PlateauTechnique.objects.filter(
            structure_id=structure_id,
        ).select_related("service")

        return Response(PlateauTechniqueSerializer(data, many=True).data)


class AdminListPlateauView(APIView):

    @admin_required
    def get(self, request):

        data = PlateauTechnique.objects.all().select_related("structure", "service")
        return Response(PlateauTechniqueSerializer(data, many=True).data)


class DeactivatePlateauView(APIView):

    @gestionnaire_required
    def patch(self, request, pk):

        plateau = PlateauTechnique.objects.select_related("structure").get(id=pk)
        assert_gestionnaire_owns_structure(request.user, plateau.structure_id)

        plateau = PlateauTechniqueService.ajouter_ou_mettre_a_jour(
            structure_id=plateau.structure_id,
            service_id=plateau.service_id,
            disponible=False,
        )

        return Response({"message": "Équipement indisponible", "data": PlateauTechniqueSerializer(plateau).data})


class ImportPlateauTechniqueView(APIView):

    @gestionnaire_required
    def post(self, request):
        structure_id = request.data.get("structure_id")
        assert_gestionnaire_owns_structure(request.user, structure_id)
        structure = Structure.objects.get(id=structure_id)

        file = request.FILES.get("file")
        if not file:
            return Response({"detail": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)

        filename = file.name.lower()
        if filename.endswith(".csv"):
            rows = self._parse_csv(file)
        elif filename.endswith((".xlsx", ".xls")):
            rows = self._parse_excel(file)
        else:
            return Response({"detail": "Format non supporté. Utilisez CSV ou Excel (.xlsx)"}, status=status.HTTP_400_BAD_REQUEST)

        if not rows:
            return Response({"detail": "Le fichier est vide ou mal formaté"}, status=status.HTTP_400_BAD_REQUEST)

        imported = []
        errors = []

        for i, row in enumerate(rows, start=2):
            try:
                nom = (row.get("nom du plateau technique") or row.get("nom plateau") or row.get("nom")).strip()
                if not nom:
                    errors.append({"ligne": i, "erreur": "Nom manquant"})
                    continue

                service, _ = StructureService.objects.get_or_create(
                    nom=nom,
                    defaults={"type": "EXAMEN", "description": ""},
                )

                obj = PlateauTechniqueService.ajouter_ou_mettre_a_jour(
                    structure_id=structure.id,
                    service_id=service.id,
                    disponible=True,
                )
                imported.append({"id": str(obj.id), "nom": service.nom})

            except Exception as e:
                errors.append({"ligne": i, "erreur": str(e)})

        return Response({
            "message": f"{len(imported)} plateau(x) technique(s) importé(s)",
            "imported": len(imported),
            "errors_count": len(errors),
            "errors": errors[:20],
        }, status=status.HTTP_200_OK)

    def _parse_csv(self, file):
        decoded = file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
        return [dict(row) for row in reader]

    def _parse_excel(self, file):
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for header, value in zip(headers, row):
                if header:
                    row_dict[header] = str(value) if value is not None else ""
            rows.append(row_dict)
        wb.close()
        return rows
