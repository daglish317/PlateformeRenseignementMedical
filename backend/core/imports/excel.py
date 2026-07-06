from openpyxl import load_workbook


class ExcelReader:
    """
    Responsable de la lecture des fichiers Excel (.xlsx)
    """

    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = None
        self.sheet = None

    def load(self):
        """
        Charge le fichier Excel en mémoire
        """
        self.workbook = load_workbook(filename=self.file_path, data_only=True)
        self.sheet = self.workbook.active
        return self

    def get_rows(self):
        """
        Retourne toutes les lignes du fichier sous forme de liste
        """

        if self.sheet is None:
            raise Exception("Fichier Excel non chargé. Appeler load() avant.")

        rows = list(self.sheet.iter_rows(values_only=True))

        if not rows:
            return []

        headers = [str(h).strip().lower() for h in rows[0]]

        data = []

        for row in rows[1:]:
            row_data = {}

            for i, value in enumerate(row):
                if i < len(headers):
                    key = headers[i]
                    row_data[key] = value

            data.append(row_data)

        return data