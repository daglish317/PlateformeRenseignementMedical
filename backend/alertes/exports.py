"""Exports PDF et Excel des alertes (module Alertes).

Seul le propriétaire peut exporter les alertes. Le gestionnaire ne
dispose pas de cette fonctionnalité.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from django.utils import timezone

_PDF_STYLE_FOND = colors.HexColor("#e2e8f0")
_PDF_STYLE_BORDURE = colors.HexColor("#cbd5e1")

_XLSX_STYLE_FOND = PatternFill(
    start_color="E2E8F0",
    end_color="E2E8F0",
    fill_type="solid",
)
_XLSX_STYLE_BORDURE = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _statut(alerte):
    return "Résolue" if alerte.est_resolue else "Active"


def _ligne_detail(alerte):
    donnees = alerte.donnees or {}
    if alerte.type == "RUPTURE_STOCK":
        return f"{donnees.get('medicament_nom', '—')} : rupture de stock"
    if alerte.type == "STOCK_FAIBLE":
        return (
            f"{donnees.get('medicament_nom', '—')} : "
            f"{donnees.get('stock_disponible', 0)} disponible(s) "
            f"pour un seuil de {donnees.get('seuil_alerte', 0)}"
        )
    if alerte.type == "RETOURS_CAISSE_ANORMAUX":
        return (
            f"{donnees.get('nombre', 0)} retour(s) caisse le "
            f"{donnees.get('date_concernee', '—')} "
            f"(moyenne {donnees.get('moyenne', 0)})"
        )
    if alerte.type == "VENTES_ANNULEES_ANORMALES":
        return (
            f"{donnees.get('nombre', 0)} annulation(s) le "
            f"{donnees.get('date_concernee', '—')} "
            f"(moyenne {donnees.get('moyenne', 0)})"
        )
    return alerte.description


def _titre_style(nom, taille=16, gras=True, align=1):
    return ParagraphStyle(
        name=nom,
        fontName="Helvetica-Bold" if gras else "Helvetica",
        fontSize=taille,
        leading=taille + 4,
        spaceAfter=2,
        alignment=align,
    )


def generer_alertes_pdf(structure, alertes, libelle_filtre=""):
    """Retourne les octets du PDF des alertes de la structure."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"Alertes - {structure.nom}",
    )

    elements = []
    elements.append(Paragraph(structure.nom or "Pharmacie", _titre_style("Titre")))
    elements.append(Paragraph(
        "Alertes détectées automatiquement",
        _titre_style("SousTitre", taille=12, gras=False),
    ))
    elements.append(Spacer(1, 6 * mm))

    infos = Table(
        [
            ["Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"],
            ["Filtre", libelle_filtre or "Toutes les alertes"],
            ["Alertes", str(alertes.count())],
        ],
        colWidths=[45 * mm, None],
    )
    infos.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(infos)
    elements.append(Spacer(1, 6 * mm))

    data = [[
        "Date",
        "Heure",
        "Priorité",
        "Catégorie",
        "Type",
        "Titre",
        "Module",
        "Statut",
    ]]
    for alerte in alertes:
        data.append([
            f"{alerte.cree_le:%d/%m/%Y}",
            f"{alerte.cree_le:%H:%M}",
            alerte.get_priorite_display(),
            alerte.get_categorie_display(),
            alerte.get_type_display(),
            alerte.titre,
            alerte.get_module_display(),
            _statut(alerte),
        ])

    tableau = Table(
        data,
        colWidths=[30 * mm, 22 * mm, 28 * mm, 35 * mm, 50 * mm, 55 * mm, 45 * mm, 22 * mm],
        repeatRows=1,
    )
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_STYLE_FOND),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_STYLE_BORDURE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(tableau)
    elements.append(Spacer(1, 8 * mm))

    elements.append(Paragraph(
        "Module Alertes : système de surveillance automatique. Aucune "
        "alerte ne peut être créée, modifiée ni supprimée par un "
        "utilisateur. Les alertes sont résolues automatiquement lorsque "
        "la situation revient à la normale.",
        _titre_style("Note", taille=8, gras=False, align=0),
    ))

    doc.build(elements)
    return buffer.getvalue()


def generer_alertes_excel(structure, alertes, libelle_filtre=""):
    """Retourne les octets du classeur Excel des alertes de la structure."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Alertes"

    ws["A1"] = structure.nom or "Pharmacie"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Alertes détectées automatiquement"

    infos = [
        ("Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"),
        ("Filtre", libelle_filtre or "Toutes les alertes"),
        ("Alertes", alertes.count()),
    ]
    for i, (cle, valeur) in enumerate(infos, start=4):
        ws.cell(row=i, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valeur)

    entete_row = 8
    entetes = [
        "Date",
        "Heure",
        "Priorité",
        "Catégorie",
        "Type",
        "Titre",
        "Module",
        "Statut",
    ]
    for col, entete in enumerate(entetes, start=1):
        cellule = ws.cell(row=entete_row, column=col, value=entete)
        cellule.font = Font(bold=True)
        cellule.fill = _XLSX_STYLE_FOND
        cellule.border = _XLSX_STYLE_BORDURE
        cellule.alignment = Alignment(horizontal="center")

    for i, alerte in enumerate(alertes, start=entete_row + 1):
        valeurs = [
            f"{alerte.cree_le:%d/%m/%Y}",
            f"{alerte.cree_le:%H:%M}",
            alerte.get_priorite_display(),
            alerte.get_categorie_display(),
            alerte.get_type_display(),
            alerte.titre,
            alerte.get_module_display(),
            _statut(alerte),
        ]
        for col, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=i, column=col, value=valeur)
            cellule.border = _XLSX_STYLE_BORDURE

    largeurs = {
        "A": 14,
        "B": 10,
        "C": 20,
        "D": 25,
        "E": 30,
        "F": 40,
        "G": 25,
        "H": 15,
    }
    for colonne, largeur in largeurs.items():
        ws.column_dimensions[colonne].width = largeur

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
