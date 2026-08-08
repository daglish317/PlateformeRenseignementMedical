"""Exports PDF et Excel du module Statistiques.

Le rapport statistique reprend la période et les filtres utilisés et
présente l'ensemble des sections d'analyse (vue générale, ventes,
produits, approvisionnements, stock, caisse, financier, comparaison).

Seul le propriétaire peut exporter les statistiques.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from django.utils import timezone

_PDF_STYLE_FOND = colors.HexColor("#e2e8f0")
_PDF_STYLE_BORDURE = colors.HexColor("#cbd5e1")

_XLSX_STYLE_FOND = PatternFill(
    start_color="E2E8F0",
    end_color="E2E8F0",
    fill_type="solid",
)
_XLSX_STYLE_BORDURE = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _montant(valeur):
    if valeur is None:
        return "—"
    return f"{valeur:,.2f}".replace(",", " ").replace(".", ",")


def _nombre(valeur):
    if valeur is None:
        return "—"
    return f"{valeur:,.0f}".replace(",", " ")


def _variation_texte(variation):
    if variation is None:
        return "—"
    if variation > 0:
        return f"+{variation:,.2f} %"
    return f"{variation:,.2f} %"


def _pourcentage(valeur):
    if valeur is None:
        return "—"
    return f"{valeur:,.2f} %"


def _titre_style(nom, taille=16, gras=True, align=1):
    return ParagraphStyle(
        name=nom,
        fontName="Helvetica-Bold" if gras else "Helvetica",
        fontSize=taille,
        leading=taille + 4,
        spaceAfter=2,
        alignment=align,
    )


def _titre_section(texte):
    return Paragraph(texte, _titre_style(f"section-{texte}", taille=12))


def _tableau_style():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_STYLE_FOND),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_STYLE_BORDURE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ])


def _infos_table(rapport):
    periode = rapport["vue_generale"]["periode"]
    libelle_periode = periode.get("libelle") or "Toute la période"
    infos = [
        ["Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"],
        ["Période", libelle_periode],
        ["Filtres", rapport["filtres"]["libelle"]],
    ]
    tableau = Table(infos, colWidths=[40 * mm, None])
    tableau.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    return tableau


def _construire_pdf_elements(structure, rapport):
    elements = []
    elements.append(Paragraph(structure.nom or "Pharmacie", _titre_style("titre")))
    elements.append(Paragraph(
        "Rapport statistique — outil de pilotage du propriétaire",
        _titre_style("sous-titre", taille=12, gras=False),
    ))
    elements.append(Spacer(1, 4 * mm))
    elements.append(_infos_table(rapport))
    elements.append(Spacer(1, 6 * mm))

    # §4 Vue générale
    vue = rapport["vue_generale"]
    elements.append(_titre_section("1. Vue générale"))
    elements.append(_tableau_donnees([
        ["Ventes validées", _nombre(vue["ventes"]["nb_validees"])],
        ["Produits vendus", _nombre(vue["ventes"]["nb_produits_vendus"])],
        [
            "Évolution des ventes",
            _variation_texte(vue["ventes"]["evolution_ventes"]),
        ],
        [
            "Évolution des produits",
            _variation_texte(vue["ventes"]["evolution_produits"]),
        ],
        ["Approvisionnements", _nombre(vue["approvisionnements"]["nombre"])],
        [
            "Quantité totale reçue",
            _nombre(vue["approvisionnements"]["quantite_totale"]),
        ],
        [
            "Évolution des approvisionnements",
            _variation_texte(vue["approvisionnements"]["evolution"]),
        ],
        ["Valeur actuelle du stock", _montant(vue["stock"]["valeur_actuelle"])],
        ["Valeur d'achat du stock", _montant(vue["stock"]["valeur_achat"])],
        ["Références disponibles", _nombre(vue["stock"]["nb_references_disponibles"])],
        ["Ruptures de stock", _nombre(vue["stock"]["nb_ruptures"])],
        ["Références sous le seuil", _nombre(vue["stock"]["nb_sous_seuil"])],
        ["Paiements validés", _nombre(vue["caisse"]["nb_paiements_valides"])],
        ["Retours caisse", _nombre(vue["caisse"]["nb_retours"])],
        ["Annulations avant paiement", _nombre(vue["caisse"]["nb_annulations"])],
    ]))
    elements.append(Spacer(1, 5 * mm))

    # §5 Analyse des ventes
    ventes = rapport["ventes"]
    elements.append(_titre_section("2. Analyse des ventes"))
    elements.append(_tableau_donnees([
        ["Ventes validées", _nombre(ventes["nb_ventes"])],
        ["Produits vendus", _nombre(ventes["nb_produits_vendus"])],
        ["Évolution des ventes", _variation_texte(ventes["evolution_ventes"])],
        [
            "Évolution des produits",
            _variation_texte(ventes["evolution_produits"]),
        ],
    ]))
    if ventes["meilleures_periodes"]:
        elements.append(Paragraph(
            "Périodes avec la plus forte activité :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [c["periode"], _nombre(c["valeur"])]
            for c in ventes["meilleures_periodes"]
        ], entetes=["Période", "Ventes"]))
    if ventes["plus_faibles_periodes"]:
        elements.append(Paragraph(
            "Périodes avec la plus faible activité :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [c["periode"], _nombre(c["valeur"])]
            for c in ventes["plus_faibles_periodes"]
        ], entetes=["Période", "Ventes"]))
    elements.append(Spacer(1, 5 * mm))

    # §6/§7 Produits
    produits = rapport["produits"]
    elements.append(_titre_section("3. Produits les plus vendus"))
    if produits["plus_vendus"]:
        elements.append(_tableau_donnees([
            [p["nom"], _nombre(p["quantite"]), _montant(p["montant"])]
            for p in produits["plus_vendus"]
        ], entetes=["Produit", "Quantité vendue", "Montant"]))
    else:
        elements.append(Paragraph(
            "Aucune vente sur la période.",
            _titre_style("aucun", taille=10, gras=False, align=0),
        ))
    elements.append(_titre_section("4. Produits les moins vendus"))
    if produits["moins_vendus"]:
        elements.append(_tableau_donnees([
            [p["nom"], _nombre(p["quantite"])]
            for p in produits["moins_vendus"]
        ], entetes=["Produit", "Quantité vendue"]))
    else:
        elements.append(Paragraph(
            "Aucun produit vendu sur la période.",
            _titre_style("aucun", taille=10, gras=False, align=0),
        ))
    elements.append(Spacer(1, 5 * mm))

    # §8 Approvisionnements
    appros = rapport["approvisionnements"]
    elements.append(_titre_section("5. Analyse des approvisionnements"))
    elements.append(_tableau_donnees([
        ["Approvisionnements", _nombre(appros["nombre"])],
        ["Quantité totale reçue", _nombre(appros["quantite_totale"])],
        ["Coût d'achat total", _montant(appros["montant_total"])],
        ["Évolution", _variation_texte(appros["evolution"])],
    ]))
    if appros["produits_frequents"]:
        elements.append(Paragraph(
            "Produits les plus réapprovisionnés :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [
                p["nom"],
                _nombre(p["nb_approvisionnements"]),
                _nombre(p["quantite_totale"]),
            ]
            for p in appros["produits_frequents"]
        ], entetes=["Produit", "Appros", "Quantité reçue"]))
    elements.append(Spacer(1, 5 * mm))

    # §9/§13 Stock
    stock = rapport["stock"]
    elements.append(_titre_section("6. Analyse du stock"))
    elements.append(_tableau_donnees([
        ["Références", _nombre(stock["nb_references"])],
        ["Références disponibles", _nombre(stock["nb_disponibles"])],
        ["Ruptures", _nombre(stock["nb_ruptures"])],
        ["Stocks faibles", _nombre(stock["nb_stocks_faibles"])],
        [
            "Évolution des ruptures",
            _variation_texte(
                (stock["evolution_ruptures"] or {}).get("variation")
            ),
        ],
    ]))
    elements.append(_titre_section("7. Valeur du stock"))
    elements.append(_tableau_donnees([
        ["Valeur d'achat", _montant(stock["valeur"]["achat"])],
        [
            "Valeur potentielle de vente",
            _montant(stock["valeur"]["vente"]),
        ],
    ]))
    if stock["produits_ruptures"]:
        elements.append(Paragraph(
            "Produits connaissant le plus de ruptures :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [p["nom"], _nombre(p["nb_ruptures"])]
            for p in stock["produits_ruptures"]
        ], entetes=["Produit", "Ruptures"]))
    elements.append(Spacer(1, 5 * mm))

    # §10/§11 Caisse
    caisse = rapport["caisse"]
    elements.append(_titre_section("8. Analyse de la caisse"))
    elements.append(_tableau_donnees([
        ["Paiements validés", _nombre(caisse["paiements"]["nombre"])],
        [
            "Montant encaissé",
            _montant(caisse["paiements"]["montant"]),
        ],
        [
            "Évolution des paiements",
            _variation_texte(caisse["paiements"]["evolution"]),
        ],
        ["Retours caisse", _nombre(caisse["retours"]["nombre"])],
        ["Montant des retours", _montant(caisse["retours"]["montant"])],
        [
            "Fréquence des retours (par jour)",
            f"{caisse['retours']['frequence_jour']:,.2f}".replace(",", "."),
        ],
        [
            "Annulations avant paiement",
            _nombre(caisse["annulations"]["nombre"]),
        ],
        [
            "Évolution des annulations",
            _variation_texte(caisse["annulations"]["evolution"]),
        ],
    ]))
    if caisse["paiements_par_mode"]:
        elements.append(Paragraph(
            "Répartition par mode de paiement :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [p["label"], _nombre(p["nombre"]), _montant(p["montant"])]
            for p in caisse["paiements_par_mode"]
        ], entetes=["Mode", "Paiements", "Montant"]))
    elements.append(Spacer(1, 5 * mm))

    # §12 Financier
    fin = rapport["financier"]
    elements.append(_titre_section("9. Analyse financière"))
    elements.append(Paragraph(
        "Information strictement réservée au propriétaire.",
        _titre_style("note", taille=8, gras=False, align=0),
    ))
    elements.append(_tableau_donnees([
        ["Chiffre d'affaires brut", _montant(fin["chiffre_affaires"]["brut"])],
        ["Retours caisse", _montant(fin["chiffre_affaires"]["retours"])],
        ["Chiffre d'affaires net", _montant(fin["chiffre_affaires"]["net"])],
        ["Ventes encaissées", _nombre(fin["nb_ventes_encaissees"])],
        ["Produits vendus", _nombre(fin["nb_produits_vendus"])],
        ["Panier moyen", _montant(fin["panier_moyen"])],
        ["Coût des marchandises vendues", _montant(fin["cout_marchandises"])],
        ["Bénéfice brut", _montant(fin["benefice_brut"])],
        ["Marge brute", _pourcentage(fin["marge_pourcentage"])],
    ]))
    if fin["par_mode"]:
        elements.append(Paragraph(
            "Répartition du chiffre d'affaires par mode :",
            _titre_style("pf", taille=10, gras=False, align=0),
        ))
        elements.append(_tableau_donnees([
            [
                p["label"],
                _montant(p["montant"]),
                _pourcentage(p["pourcentage"]),
            ]
            for p in fin["par_mode"]
        ], entetes=["Mode", "Montant", "Part"]))
    elements.append(Spacer(1, 5 * mm))

    # §14 Comparaison
    comparaison = rapport["comparaison"]
    if comparaison["periode_precedente"]:
        elements.append(_titre_section("10. Comparaison des périodes"))
        elements.append(_tableau_donnees([
            ["Période actuelle", comparaison["periode_actuelle"].get("libelle")],
            [
                "Période précédente",
                comparaison["periode_precedente"].get("libelle"),
            ],
            [
                "Ventes",
                _comparaison_cellule(comparaison["ventes"]),
            ],
            [
                "Approvisionnements",
                _comparaison_cellule(comparaison["approvisionnements"]),
            ],
            [
                "Ruptures",
                _comparaison_cellule(comparaison["ruptures"]),
            ],
            [
                "Retours caisse",
                _comparaison_cellule(comparaison["retours_caisse"]),
            ],
        ]))

    elements.append(Spacer(1, 8 * mm))
    elements.append(Paragraph(
        "Module Statistiques : couche d'analyse du système, strictement en "
        "lecture. Aucune action ne permet de modifier le stock, une vente, "
        "un approvisionnement, la caisse ou un inventaire. Les informations "
        "financières sont exclusivement réservées au propriétaire.",
        _titre_style("note-finale", taille=8, gras=False, align=0),
    ))

    return elements


def _comparaison_cellule(donnees):
    if donnees is None:
        return "—"
    precedent = donnees.get("precedent")
    actuel = donnees.get("actuel")
    return (
        f"Précédent : {_nombre(precedent)} — actuel : {_nombre(actuel)} "
        f"({_variation_texte(donnees.get('variation'))})"
    )


def _tableau_donnees(lignes, entetes=None, largeurs=None):
    donnees = [entetes] if entetes else []
    donnees.extend([[str(v) for v in ligne] for ligne in lignes])
    tableau = Table(
        donnees,
        colWidths=largeurs,
        repeatRows=1 if entetes else 0,
    )
    tableau.setStyle(_tableau_style())
    return tableau


def generer_statistiques_pdf(structure, rapport):
    """Retourne les octets du PDF du rapport statistique."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"Statistiques - {structure.nom}",
    )
    elements = _construire_pdf_elements(structure, rapport)
    doc.build(elements)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def _excel_titre(ws, structure, rapport, titre):
    ws["A1"] = structure.nom or "Pharmacie"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = titre
    periode = rapport["vue_generale"]["periode"].get("libelle") or "Toute la période"
    infos = [
        ("Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"),
        ("Période", periode),
        ("Filtres", rapport["filtres"]["libelle"]),
    ]
    for i, (cle, valeur) in enumerate(infos, start=4):
        ws.cell(row=i, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valeur)


def _excel_entete(ws, ligne, entetes):
    for col, entete in enumerate(entetes, start=1):
        cellule = ws.cell(row=ligne, column=col, value=entete)
        cellule.font = Font(bold=True)
        cellule.fill = _XLSX_STYLE_FOND
        cellule.border = _XLSX_STYLE_BORDURE
        cellule.alignment = Alignment(horizontal="center")


def _excel_ligne(ws, ligne, valeurs):
    for col, valeur in enumerate(valeurs, start=1):
        cellule = ws.cell(row=ligne, column=col, value=valeur)
        cellule.border = _XLSX_STYLE_BORDURE


def _excel_cartes(ws, ligne, paires):
    for i, (cle, valeur) in enumerate(paires, start=1):
        ws.cell(row=ligne, column=2 * i - 1, value=cle).font = Font(bold=True)
        ws.cell(row=ligne, column=2 * i, value=valeur)


def generer_statistiques_excel(structure, rapport):
    """Retourne les octets du classeur Excel du rapport statistique."""
    wb = Workbook()

    # Feuille 1 : Vue générale
    ws = wb.active
    ws.title = "Vue generale"
    _excel_titre(ws, structure, rapport, "Vue générale de l'activité")
    vue = rapport["vue_generale"]
    ligne = 9
    ws.cell(row=ligne, column=1, value="VENTES").font = Font(bold=True)
    ligne += 1
    for cle, valeur in [
        ("Ventes validées", _nombre(vue["ventes"]["nb_validees"])),
        ("Produits vendus", _nombre(vue["ventes"]["nb_produits_vendus"])),
        (
            "Évolution des ventes",
            _variation_texte(vue["ventes"]["evolution_ventes"]),
        ),
    ]:
        ws.cell(row=ligne, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1
    ws.cell(row=ligne, column=1, value="APPROVISIONNEMENTS").font = Font(bold=True)
    ligne += 1
    for cle, valeur in [
        ("Approvisionnements", _nombre(vue["approvisionnements"]["nombre"])),
        (
            "Quantité totale reçue",
            _nombre(vue["approvisionnements"]["quantite_totale"]),
        ),
        (
            "Évolution",
            _variation_texte(vue["approvisionnements"]["evolution"]),
        ),
    ]:
        ws.cell(row=ligne, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1
    ws.cell(row=ligne, column=1, value="STOCK").font = Font(bold=True)
    ligne += 1
    for cle, valeur in [
        ("Valeur actuelle du stock", _montant(vue["stock"]["valeur_actuelle"])),
        ("Valeur d'achat", _montant(vue["stock"]["valeur_achat"])),
        ("Références disponibles", _nombre(vue["stock"]["nb_references_disponibles"])),
        ("Ruptures", _nombre(vue["stock"]["nb_ruptures"])),
        ("Sous le seuil", _nombre(vue["stock"]["nb_sous_seuil"])),
    ]:
        ws.cell(row=ligne, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1
    ws.cell(row=ligne, column=1, value="CAISSE").font = Font(bold=True)
    ligne += 1
    for cle, valeur in [
        ("Paiements validés", _nombre(vue["caisse"]["nb_paiements_valides"])),
        ("Retours caisse", _nombre(vue["caisse"]["nb_retours"])),
        (
            "Montant des retours",
            _montant(vue["caisse"]["montant_retours"]),
        ),
        ("Annulations avant paiement", _nombre(vue["caisse"]["nb_annulations"])),
    ]:
        ws.cell(row=ligne, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=ligne, column=2, value=valeur)
        ligne += 1
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 22

    # Feuille 2 : Ventes
    ws2 = wb.create_sheet("Ventes")
    _excel_titre(ws2, structure, rapport, "Analyse des ventes")
    ventes = rapport["ventes"]
    _excel_cartes(ws2, 8, [
        ("Ventes validées", ventes["nb_ventes"]),
        ("Produits vendus", ventes["nb_produits_vendus"]),
        ("Évolution", _variation_texte(ventes["evolution_ventes"])),
    ])
    _excel_entete(ws2, 10, ["Période", "Ventes", "Produits vendus"])
    for i, (sv, sp) in enumerate(
        zip(ventes["series_ventes"], ventes["series_produits"]), start=11
    ):
        _excel_ligne(ws2, i, [sv["periode"], sv["valeur"], sp["valeur"]])

    # Feuille 3 : Produits
    ws3 = wb.create_sheet("Produits")
    _excel_titre(ws3, structure, rapport, "Produits les plus / les moins vendus")
    produits = rapport["produits"]
    ws3.cell(row=8, column=1, value="PLUS VENDUS").font = Font(bold=True)
    _excel_entete(ws3, 9, ["Produit", "Quantité vendue", "Montant"])
    for i, p in enumerate(produits["plus_vendus"], start=10):
        _excel_ligne(ws3, i, [p["nom"], p["quantite"], p["montant"]])
    ligne3 = 10 + len(produits["plus_vendus"]) + 2
    ws3.cell(row=ligne3, column=1, value="MOINS VENDUS").font = Font(bold=True)
    _excel_entete(ws3, ligne3 + 1, ["Produit", "Quantité vendue"])
    for i, p in enumerate(produits["moins_vendus"], start=ligne3 + 2):
        _excel_ligne(ws3, i, [p["nom"], p["quantite"]])

    # Feuille 4 : Approvisionnements
    ws4 = wb.create_sheet("Approvisionnements")
    _excel_titre(ws4, structure, rapport, "Analyse des approvisionnements")
    appros = rapport["approvisionnements"]
    _excel_cartes(ws4, 8, [
        ("Nombre", appros["nombre"]),
        ("Quantité reçue", appros["quantite_totale"]),
        ("Coût total", _montant(appros["montant_total"])),
        ("Évolution", _variation_texte(appros["evolution"])),
    ])
    _excel_entete(ws4, 10, ["Période", "Approvisionnements", "Quantité"])
    for i, (sn, sq) in enumerate(
        zip(appros["series_nombre"], appros["series_quantite"]), start=11
    ):
        _excel_ligne(ws4, i, [sn["periode"], sn["valeur"], sq["valeur"]])
    ligne4 = 11 + max(len(appros["series_nombre"]), 1) + 2
    ws4.cell(row=ligne4, column=1, value="PRODUITS FREQUENTS").font = Font(bold=True)
    _excel_entete(ws4, ligne4 + 1, ["Produit", "Appros", "Quantité reçue"])
    for i, p in enumerate(appros["produits_frequents"], start=ligne4 + 2):
        _excel_ligne(ws4, i, [
            p["nom"],
            p["nb_approvisionnements"],
            p["quantite_totale"],
        ])

    # Feuille 5 : Stock
    ws5 = wb.create_sheet("Stock")
    _excel_titre(ws5, structure, rapport, "Analyse du stock")
    stock = rapport["stock"]
    _excel_cartes(ws5, 8, [
        ("Références", stock["nb_references"]),
        ("Disponibles", stock["nb_disponibles"]),
        ("Ruptures", stock["nb_ruptures"]),
        ("Stocks faibles", stock["nb_stocks_faibles"]),
    ])
    evo_ruptures = stock.get("evolution_ruptures") or {}
    _excel_cartes(ws5, 10, [
        ("Évolution des ruptures", _variation_texte(evo_ruptures.get("variation"))),
        ("Valeur d'achat", _montant(stock["valeur"]["achat"])),
        ("Valeur de vente", _montant(stock["valeur"]["vente"])),
        ("", ""),
    ])
    _excel_entete(ws5, 12, ["Type", "Références", "Disponibles", "Ruptures", "Faibles"])
    for i, pt in enumerate(stock["par_type"], start=13):
        _excel_ligne(ws5, i, [
            pt["type"],
            pt["nb_references"],
            pt["nb_disponibles"],
            pt["nb_ruptures"],
            pt["nb_stocks_faibles"],
        ])
    ligne5 = 13 + len(stock["par_type"]) + 2
    ws5.cell(row=ligne5, column=1, value="PRODUITS EN RUPTURE").font = Font(bold=True)
    _excel_entete(ws5, ligne5 + 1, ["Produit", "Ruptures"])
    for i, p in enumerate(stock["produits_ruptures"], start=ligne5 + 2):
        _excel_ligne(ws5, i, [p["nom"], p["nb_ruptures"]])

    # Feuille 6 : Caisse
    ws6 = wb.create_sheet("Caisse")
    _excel_titre(ws6, structure, rapport, "Analyse de la caisse")
    caisse = rapport["caisse"]
    _excel_cartes(ws6, 8, [
        ("Paiements validés", caisse["paiements"]["nombre"]),
        ("Montant encaissé", _montant(caisse["paiements"]["montant"])),
        ("Retours caisse", caisse["retours"]["nombre"]),
        ("Montant des retours", _montant(caisse["retours"]["montant"])),
    ])
    _excel_cartes(ws6, 10, [
        ("Annulations avant paiement", caisse["annulations"]["nombre"]),
        ("Évolution retours", _variation_texte(caisse["retours"]["evolution"])),
        ("Évolution annulations", _variation_texte(caisse["annulations"]["evolution"])),
        ("Fréquence retours/jour", caisse["retours"]["frequence_jour"]),
    ])
    _excel_entete(ws6, 12, ["Mode de paiement", "Nombre", "Montant"])
    for i, p in enumerate(caisse["paiements_par_mode"], start=13):
        _excel_ligne(ws6, i, [p["label"], p["nombre"], p["montant"]])

    # Feuille 7 : Financier
    ws7 = wb.create_sheet("Financier")
    _excel_titre(ws7, structure, rapport, "Analyse financière (propriétaire)")
    fin = rapport["financier"]
    _excel_cartes(ws7, 8, [
        ("CA brut", _montant(fin["chiffre_affaires"]["brut"])),
        ("Retours", _montant(fin["chiffre_affaires"]["retours"])),
        ("CA net", _montant(fin["chiffre_affaires"]["net"])),
        ("Panier moyen", _montant(fin["panier_moyen"])),
    ])
    _excel_cartes(ws7, 10, [
        ("Coût des marchandises", _montant(fin["cout_marchandises"])),
        ("Bénéfice brut", _montant(fin["benefice_brut"])),
        ("Marge brute", _pourcentage(fin["marge_pourcentage"])),
        ("", ""),
    ])
    _excel_entete(ws7, 12, ["Mode de paiement", "Nombre", "Montant", "Part"])
    for i, p in enumerate(fin["par_mode"], start=13):
        _excel_ligne(ws7, i, [
            p["label"],
            p["nombre"],
            p["montant"],
            _pourcentage(p["pourcentage"]),
        ])

    # Feuille 8 : Comparaison
    ws8 = wb.create_sheet("Comparaison")
    _excel_titre(ws8, structure, rapport, "Comparaison des périodes")
    comp = rapport["comparaison"]
    _excel_entete(ws8, 8, ["Indicateur", "Période précédente", "Période actuelle", "Variation"])
    if comp["periode_precedente"]:
        for i, (cle, donnees) in enumerate([
            ("Ventes", comp["ventes"]),
            ("Approvisionnements", comp["approvisionnements"]),
            ("Ruptures", comp["ruptures"]),
            ("Retours caisse", comp["retours_caisse"]),
        ], start=9):
            if donnees is None:
                continue
            _excel_ligne(ws8, i, [
                cle,
                donnees.get("precedent"),
                donnees.get("actuel"),
                _variation_texte(donnees.get("variation")),
            ])

    # Feuille 9 : Détails
    ws9 = wb.create_sheet("Details")
    _excel_titre(ws9, structure, rapport, "Détails statistiques")
    details = rapport["details"]
    ligne9 = 8
    ws9.cell(row=ligne9, column=1, value="VENTES VALIDEES").font = Font(bold=True)
    _excel_entete(ws9, ligne9 + 1, ["Numéro", "Date", "Montant", "Articles", "Mode"])
    ligne9 += 2
    for v in details["ventes"]["items"]:
        _excel_ligne(ws9, ligne9, [
            v["numero"],
            v["validee_le"],
            v["montant_total"],
            v["nb_articles"],
            v["mode_label"],
        ])
        ligne9 += 1
    ligne9 += 1
    ws9.cell(row=ligne9, column=1, value="APPROVISIONNEMENTS").font = Font(bold=True)
    _excel_entete(ws9, ligne9 + 1, ["Numéro", "Date", "Fournisseur", "Quantité", "Montant"])
    ligne9 += 2
    for a in details["approvisionnements"]["items"]:
        _excel_ligne(ws9, ligne9, [
            a["numero"],
            a["date_reception"],
            a["fournisseur"],
            a["quantite_totale"],
            a["montant_total"],
        ])
        ligne9 += 1
    ligne9 += 1
    ws9.cell(row=ligne9, column=1, value="RETOURS CAISSE").font = Font(bold=True)
    _excel_entete(ws9, ligne9 + 1, ["Numéro", "Date", "Motif", "Montant", "Articles"])
    ligne9 += 2
    for r in details["retours"]["items"]:
        _excel_ligne(ws9, ligne9, [
            r["numero"],
            r["effectue_le"],
            r["motif"],
            r["montant_total"],
            r["nb_articles"],
        ])
        ligne9 += 1
    ligne9 += 1
    ws9.cell(row=ligne9, column=1, value="ANNULATIONS AVANT PAIEMENT").font = Font(bold=True)
    _excel_entete(ws9, ligne9 + 1, ["Numéro", "Date", "Motif d'annulation"])
    ligne9 += 2
    for ann in details["annulations"]["items"]:
        _excel_ligne(ws9, ligne9, [
            ann["numero"],
            ann["annulee_le"],
            ann["motif_annulation"],
        ])
        ligne9 += 1

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
