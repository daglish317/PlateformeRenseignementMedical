import csv
import io

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from .models import PriseEnCharge
from .serializers import PriseEnChargeSerializer, PriseEnChargeCreateSerializer, PriseEnChargeUpdateSerializer
from .services import PriseEnChargeService

from structures.models import Structure, StructureService
from utilisateurs.decorators import gestionnaire_required, admin_required
from structures.permissions import assert_gestionnaire_owns_structure


class CreatePriseEnChargeView(APIView):

    @gestionnaire_required
    def post(self, request):

        serializer = PriseEnChargeCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        assert_gestionnaire_owns_structure(request.user, data["structure_id"])

        if not data.get("service_id") and data.get("nom"):
            service_obj, _ = StructureService.objects.get_or_create(
                nom=data["nom"],
                defaults={"type": "SERVICE_MEDICAL", "description": ""},
            )
            data["service_id"] = service_obj.id

        prise = PriseEnChargeService.ajouter_ou_mettre_a_jour(
            structure_id=data["structure_id"],
            service_id=data["service_id"],
            niveau=data.get("niveau"),
        )

        return Response(
            {"message": "Prise en charge enregistrée", "data": PriseEnChargeSerializer(prise).data},
            status=status.HTTP_201_CREATED,
        )


class UpdatePriseEnChargeView(APIView):

    @gestionnaire_required
    def patch(self, request, pk):

        try:
            prise_en_charge = PriseEnCharge.objects.select_related("structure").get(id=pk)
        except PriseEnCharge.DoesNotExist:
            return Response(
                {"detail": "Prise en charge introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_gestionnaire_owns_structure(request.user, prise_en_charge.structure_id)

        serializer = PriseEnChargeUpdateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        prise_en_charge = PriseEnChargeService.modifier(
            prise_en_charge=prise_en_charge,
            niveau=serializer.validated_data.get("niveau"),
        )

        return Response(
            {"message": "Prise en charge mise à jour", "data": PriseEnChargeSerializer(prise_en_charge).data},
            status=status.HTTP_200_OK,
        )


class DeletePriseEnChargeView(APIView):

    @gestionnaire_required
    def delete(self, request, pk):

        try:
            prise_en_charge = PriseEnCharge.objects.select_related("structure").get(id=pk)
        except PriseEnCharge.DoesNotExist:
            return Response(
                {"detail": "Prise en charge introuvable"},
                status=status.HTTP_404_NOT_FOUND,
            )

        assert_gestionnaire_owns_structure(request.user, prise_en_charge.structure_id)

        PriseEnChargeService.supprimer(prise_en_charge=prise_en_charge)

        return Response(
            {"message": "Prise en charge supprimée"},
            status=status.HTTP_200_OK,
        )


class ListPriseEnChargeStructureView(APIView):

    def get(self, request, structure_id):

        prises = PriseEnCharge.objects.filter(
            structure_id=structure_id,
        ).select_related("service")

        niveau = request.query_params.get("niveau")
        type_service = request.query_params.get("type")

        if niveau:
            prises = prises.filter(niveau=niveau)
        if type_service:
            prises = prises.filter(service__type=type_service)

        return Response(PriseEnChargeSerializer(prises, many=True).data)


class AdminListPriseEnChargeView(APIView):

    @admin_required
    def get(self, request):

        prises = PriseEnCharge.objects.all().select_related("structure", "service")
        return Response(PriseEnChargeSerializer(prises, many=True).data)


class ImportPriseEnChargeView(APIView):

    @gestionnaire_required
    def post(self, request):
        structure_id = request.data.get("structure_id")
        assert_gestionnaire_owns_structure(request.user, structure_id)
        structure = Structure.objects.get(id=structure_id)

        file = request.FILES.get("file")
        if not file:
            return Response({"detail": "Aucun fichier fourni"}, status=status.HTTP_400_BAD_REQUEST)

        filename = file.name.lower()
        if filename.endswith(".csv"):
            rows = self._parse_csv(file)
        elif filename.endswith((".xlsx", ".xls")):
            rows = self._parse_excel(file)
        else:
            return Response({"detail": "Format non supporté. Utilisez CSV ou Excel (.xlsx)"}, status=status.HTTP_400_BAD_REQUEST)

        if not rows:
            return Response({"detail": "Le fichier est vide ou mal formaté"}, status=status.HTTP_400_BAD_REQUEST)

        imported = []
        errors = []

        for i, row in enumerate(rows, start=2):
            try:
                nom = (row.get("nom de la prise en charge") or row.get("nom prise en charge") or row.get("nom")).strip()
                if not nom:
                    errors.append({"ligne": i, "erreur": "Nom manquant"})
                    continue

                service, _ = StructureService.objects.get_or_create(
                    nom=nom,
                    defaults={"type": "SERVICE_MEDICAL", "description": ""},
                )

                obj = PriseEnChargeService.ajouter_ou_mettre_a_jour(
                    structure_id=structure.id,
                    service_id=service.id,
                    niveau=None,
                )
                imported.append({"id": str(obj.id), "nom": service.nom})

            except Exception as e:
                errors.append({"ligne": i, "erreur": str(e)})

        return Response({
            "message": f"{len(imported)} prise(s) en charge importée(s)",
            "imported": len(imported),
            "errors_count": len(errors),
            "errors": errors[:20],
        }, status=status.HTTP_200_OK)

    def _parse_csv(self, file):
        decoded = file.read().decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
        return [dict(row) for row in reader]

    def _parse_excel(self, file):
        import openpyxl
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for header, value in zip(headers, row):
                if header:
                    row_dict[header] = str(value) if value is not None else ""
            rows.append(row_dict)
        wb.close()
        return rows
