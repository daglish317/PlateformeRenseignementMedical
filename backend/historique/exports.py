"""Exports PDF et Excel de l'historique (module Historique).

Seul le propriétaire peut exporter l'historique. Le gestionnaire ne
dispose pas de cette fonctionnalité.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from django.utils import timezone

_PDF_STYLE_FOND = colors.HexColor("#e2e8f0")
_PDF_STYLE_BORDURE = colors.HexColor("#cbd5e1")

_XLSX_STYLE_FOND = PatternFill(
    start_color="E2E8F0",
    end_color="E2E8F0",
    fill_type="solid",
)
_XLSX_STYLE_BORDURE = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)


def _ligne_resume(evenement):
    donnees = evenement.donnees or {}
    numero = donnees.get("numero") or donnees.get("numero_retour") or "—"
    if evenement.type == "CAISSE_RETOUR":
        detail = f"Retour {numero} — vente {donnees.get('numero_vente', '—')} / facture {donnees.get('numero_facture', '—')} — {donnees.get('nb_articles', 0)} article(s)"
    elif evenement.type == "INVENTAIRE_GENERE":
        detail = f"Inventaire {numero} — {donnees.get('nb_produits', 0)} produit(s)"
    else:
        detail = f"Approvisionnement {numero} — {donnees.get('nb_produits', 0)} produit(s)"
        if donnees.get("fournisseur"):
            detail += f" — {donnees['fournisseur']}"
    return detail


def _titre_style(nom, taille=16, gras=True, align=1):
    return ParagraphStyle(
        name=nom,
        fontName="Helvetica-Bold" if gras else "Helvetica",
        fontSize=taille,
        leading=taille + 4,
        spaceAfter=2,
        alignment=align,
    )


def generer_historique_pdf(structure, evenements, periode=""):
    """Retourne les octets du PDF du journal d'historique."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=f"Historique - {structure.nom}",
    )

    elements = []
    elements.append(Paragraph(structure.nom or "Pharmacie", _titre_style("Titre")))
    elements.append(Paragraph(
        "Historique des événements majeurs",
        _titre_style("SousTitre", taille=12, gras=False),
    ))
    elements.append(Spacer(1, 6 * mm))

    infos = Table(
        [
            ["Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"],
            ["Période", periode or "Toute la période"],
            ["Événements", str(evenements.count())],
        ],
        colWidths=[45 * mm, None],
    )
    infos.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    elements.append(infos)
    elements.append(Spacer(1, 6 * mm))

    data = [[
        "Date",
        "Heure",
        "Type d'événement",
        "Utilisateur",
        "Rôle",
        "Détail",
    ]]
    for evenement in evenements:
        data.append([
            f"{evenement.cree_le:%d/%m/%Y}",
            f"{evenement.cree_le:%H:%M}",
            evenement.get_type_display(),
            evenement.utilisateur.nom,
            evenement.role or "—",
            _ligne_resume(evenement),
        ])

    tableau = Table(
        data,
        colWidths=[30 * mm, 22 * mm, 45 * mm, 45 * mm, 30 * mm, None],
        repeatRows=1,
    )
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _PDF_STYLE_FOND),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.4, _PDF_STYLE_BORDURE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elements.append(tableau)
    elements.append(Spacer(1, 8 * mm))

    elements.append(Paragraph(
        "Journal de traçabilité : les événements sont classés du plus "
        "récent au plus ancien. Aucun événement ne peut être modifié ni "
        "supprimé. Les opérations courantes restent consultables dans "
        "leurs modules respectifs.",
        _titre_style("Note", taille=8, gras=False, align=0),
    ))

    doc.build(elements)
    return buffer.getvalue()


def generer_historique_excel(structure, evenements, periode=""):
    """Retourne les octets du classeur Excel du journal d'historique."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Historique"

    ws["A1"] = structure.nom or "Pharmacie"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Historique des événements majeurs"

    infos = [
        ("Édité le", f"{timezone.now():%d/%m/%Y %H:%M}"),
        ("Période", periode or "Toute la période"),
        ("Événements", evenements.count()),
    ]
    for i, (cle, valeur) in enumerate(infos, start=4):
        ws.cell(row=i, column=1, value=cle).font = Font(bold=True)
        ws.cell(row=i, column=2, value=valeur)

    entete_row = 8
    entetes = [
        "Date",
        "Heure",
        "Type d'événement",
        "Utilisateur",
        "Rôle",
        "Détail",
    ]
    for col, entete in enumerate(entetes, start=1):
        cellule = ws.cell(row=entete_row, column=col, value=entete)
        cellule.font = Font(bold=True)
        cellule.fill = _XLSX_STYLE_FOND
        cellule.border = _XLSX_STYLE_BORDURE
        cellule.alignment = Alignment(horizontal="center")

    for i, evenement in enumerate(evenements, start=entete_row + 1):
        valeurs = [
            f"{evenement.cree_le:%d/%m/%Y}",
            f"{evenement.cree_le:%H:%M}",
            evenement.get_type_display(),
            evenement.utilisateur.nom,
            evenement.role or "—",
            _ligne_resume(evenement),
        ]
        for col, valeur in enumerate(valeurs, start=1):
            cellule = ws.cell(row=i, column=col, value=valeur)
            cellule.border = _XLSX_STYLE_BORDURE

    largeurs = {
        "A": 14,
        "B": 10,
        "C": 35,
        "D": 30,
        "E": 20,
        "F": 70,
    }
    for colonne, largeur in largeurs.items():
        ws.column_dimensions[colonne].width = largeur

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
